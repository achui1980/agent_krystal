"""
示例脚本：使用 utils 工具库进行数据转换

这个脚本演示了如何使用 krystal_v2.case_generator.utils 工具库
来读取源文件、转换数据并生成输出文件。

运行方式:
    cd /Users/portz/js/agent-krystal
    python3 krystal_v2/case_generator/utils/example_usage.py
"""

import sys
import importlib.util
from pathlib import Path


# 由于 krystal_v2/__init__.py 需要 crewai，我们直接加载模块文件
def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


# 加载所需模块
base_path = Path(__file__).parent
shared_rules = load_module("shared_rules", base_path / "shared_rules.py")
file_readers = load_module("file_readers", base_path / "file_readers.py")
data_converters = load_module("data_converters", base_path / "data_converters.py")
formatters = load_module("formatters", base_path / "formatters.py")


def transform_humana_s10_data():
    """
    转换 humanaS10 案例的数据
    """
    print("=" * 80)
    print("示例：转换 humanaS10 案例数据")
    print("=" * 80)

    # 1. 读取源文件
    print("\n[1] 读取源文件...")
    source_df = file_readers.read_source_file("case/humanaS10/source.csv")
    print(f"    ✅ 读取成功：{len(source_df)} 行, {len(source_df.columns)} 列")

    # 2. 读取 rules.xlsx 获取案例特定的固定值（如 CARRIER_FAMILY_ID）
    print("\n[2] 读取 rules.xlsx...")

    # 直接使用 openpyxl 读取 rules.xlsx（因为结构特殊，第6行是表头，第7行开始是数据）
    import openpyxl

    wb = openpyxl.load_workbook("case/humanaS10/rules.xlsx", data_only=True)
    ws = wb["Sheet1"]

    # 提取 CARRIER_FAMILY_ID 的值（在第8行，列D）
    # 注意：CARRIER_FAMILY_ID 格式为字符串（如 "66,175,206" 或 "66"）
    carrier_family_id = None
    for row in ws.iter_rows(min_row=7, max_row=100):
        cs_col_name = row[1].value if len(row) > 1 else None
        default_value = row[3].value if len(row) > 3 else None

        if cs_col_name == "CARRIER_FAMILY_ID":
            # 直接作为字符串读取，保持原始格式
            if default_value is not None:
                carrier_family_id = str(default_value).strip()
            else:
                carrier_family_id = ""
            break

        if not cs_col_name:
            break

    print(f"    ✅ CARRIER_FAMILY_ID: {carrier_family_id}")

    # 3. 提取字段映射规则（从 rules.xlsx）
    print("\n[3] 提取字段映射规则...")
    field_mappings = {}
    special_rules = {}

    for row in ws.iter_rows(min_row=7, max_row=100):
        cs_col_name = row[1].value if len(row) > 1 else None
        carrier_col_name = row[2].value if len(row) > 2 else None  # CARRIER_COLUMN_NAME
        special_rule = row[4].value if len(row) > 4 else None  # SPECIAL_RULES

        if not cs_col_name:
            break

        # 跳过空字段和固定值字段
        if (
            carrier_col_name
            and carrier_col_name not in ["NULL", "n/a"]
            and not isinstance(carrier_col_name, (int, float))
        ):
            field_mappings[cs_col_name] = str(carrier_col_name)

            # 保存特殊规则
            if special_rule:
                special_rules[cs_col_name] = str(special_rule)

    print(f"    ✅ 找到 {len(field_mappings)} 个字段映射")
    print(f"    ✅ 找到 {len(special_rules)} 个特殊规则")

    # 4. 转换数据
    print("\n[4] 转换数据...")
    output_rows = []

    for idx, row in source_df.iterrows():
        # 过滤条件：跳过 Eff_Date == Term_Date 的行（无效数据）
        eff_date = data_converters.safe_get(row, "Eff_Date", "")
        term_date = data_converters.safe_get(row, "Term_Date", "")

        if eff_date and term_date and eff_date == term_date:
            # 跳过生效日期和终止日期相同的无效记录
            continue

        output_row = {}

        # 应用共享固定值
        output_row.update(shared_rules.SHARED_FIXED_VALUES)

        # 应用案例特定固定值
        output_row["CARRIER_FAMILY_ID"] = carrier_family_id

        # 应用空字段
        for field in shared_rules.SHARED_EMPTY_FIELDS:
            output_row[field] = ""

        # 从源数据映射字段（基于 field_mappings）
        for target_field, source_field in field_mappings.items():
            original_value = data_converters.safe_get(row, source_field, "")
            value = original_value

            # 处理特殊规则
            if target_field in special_rules:
                rule = special_rules[target_field]

                # 姓名解析规则: "Member format: last_name, first_name"
                if "last_name, first_name" in rule.lower() and value:
                    parsed = data_converters.parse_full_name(value)
                    if target_field == "FIRST_NAME":
                        # FIRST_NAME 包含名字和中间名
                        if parsed["middle"]:
                            value = f"{parsed['first']} {parsed['middle']}"
                        else:
                            value = parsed["first"]
                    elif target_field == "LAST_NAME":
                        value = parsed["last"]
                    elif target_field == "MIDDLE_INITIAL":
                        # MIDDLE_INITIAL 保持为空（已在 FIRST_NAME 中）
                        value = ""

                # CMS合同分割规则: "Plan_Name" contains "CMS_Contract_ID-CMS_Plan_CODE"
                # 但只对 MD 和 MA/MAPD 产品类型有效，MS 类型不分割
                elif (
                    "cms_contract_id" in rule.lower()
                    and "cms_plan_code" in rule.lower()
                    and value
                ):
                    product = data_converters.safe_get(row, "Product", "")
                    product_line = data_converters.map_product_type(product)

                    # 只有 MD 和 MA/MAPD 需要分割
                    if product_line in ["MD", "MA/MAPD"]:
                        cms_split = data_converters.split_cms_contract(value)
                        if target_field == "CMS_CONTRACT_ID":
                            value = cms_split["contract"]
                        elif target_field == "CMS_PLAN_ID":
                            value = cms_split["plan"]
                    else:
                        # MS, HV, 等其他类型不分割，保持为空
                        value = ""

            # 特殊处理：CMS_PLAN_ID 也需要从 Plan_Name 分割（即使没有 special_rules）
            # 但只对 MD 和 MA/MAPD 产品类型有效
            elif (
                target_field == "CMS_PLAN_ID"
                and source_field == "Plan_Name"
                and original_value
            ):
                product = data_converters.safe_get(row, "Product", "")
                product_line = data_converters.map_product_type(product)

                if product_line in ["MD", "MA/MAPD"]:
                    cms_split = data_converters.split_cms_contract(original_value)
                    value = cms_split["plan"]
                else:
                    value = ""

            # 特殊处理：日期字段
            if "DATE" in target_field and value:
                value = data_converters.convert_date_to_mmddyyyy(value)

            # 特殊处理：产品类型映射
            if target_field == "PRODUCT_LINE" and value:
                value = data_converters.map_product_type(value)

            output_row[target_field] = value

        # 特殊处理：CARRIER_STATUS_MAP (基于 Term_Date)
        term_date = data_converters.safe_get(row, "Term_Date", "")

        # 如果 Term_Date 不存在或为 nan，视为 Active
        if not term_date or term_date == "nan" or term_date == "9999-12-31":
            output_row["CARRIER_STATUS_MAP"] = "Active"
        else:
            # 有具体的终止日期（且不是永久），状态为 Termed
            output_row["CARRIER_STATUS_MAP"] = "Termed"

        # LAST_TOUCHED_DATE 设为空
        output_row["LAST_TOUCHED_DATE"] = ""

        output_rows.append(output_row)

    print(f"    ✅ 转换完成：{len(output_rows)} 行")

    # 5. 生成输出
    print("\n[5] 生成输出文件...")
    metadata = formatters.create_metadata(
        action_id="humana-s10-cs-data-integration", service_map_id="10003358"
    )

    output_content = formatters.format_pipe_delimited_output(
        output_rows, shared_rules.EXPECTED_FIELD_ORDER, metadata
    )

    # 显示输出预览
    lines = output_content.split("\n")
    print(f"    ✅ 生成成功：共 {len(lines)} 行")
    print("\n    输出预览（前10行）:")
    for i, line in enumerate(lines[:10], 1):
        if len(line) > 120:
            print(f"      {i:2d}: {line[:120]}...")
        else:
            print(f"      {i:2d}: {line}")

    # 6. 保存到文件（可选）
    output_file = "tmp/example_output.txt"
    Path("tmp").mkdir(exist_ok=True)
    with open(output_file, "w", encoding="utf-8") as f:
        f.write(output_content)
    print(f"\n    💾 输出已保存到: {output_file}")

    print("\n" + "=" * 80)
    print("✅ 示例完成！")
    print("=" * 80)


def demo_converter_functions():
    """
    演示各种转换函数的用法
    """
    print("\n" + "=" * 80)
    print("示例：数据转换函数演示")
    print("=" * 80)

    print("\n[日期转换]")
    print(f"  Excel 44197 → {data_converters.convert_date_to_mmddyyyy(44197)}")
    print(
        f"  ISO 2021-01-01 → {data_converters.convert_date_to_mmddyyyy('2021-01-01')}"
    )

    print("\n[姓名解析]")
    name = data_converters.parse_full_name("SMITH,JOHN MICHAEL")
    print(
        f"  'SMITH,JOHN MICHAEL' → first={name['first']}, middle={name['middle']}, last={name['last']}"
    )

    print("\n[CMS合同拆分]")
    cms = data_converters.split_cms_contract("S5884-197")
    print(f"  'S5884-197' → contract={cms['contract']}, plan={cms['plan']}")

    print("\n[数字格式化]")
    print(f"  66175206 → {data_converters.format_with_commas(66175206)}")

    print("\n" + "=" * 80)


if __name__ == "__main__":
    # 运行示例
    demo_converter_functions()
    print("\n")
    transform_humana_s10_data()
