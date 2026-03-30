"""
规则解析器

读取 rules.xlsx 并解析为结构化配置
"""

import openpyxl
from pathlib import Path
from typing import List, Dict, Any, Optional
import sys

# 添加父目录到路径
sys.path.insert(0, str(Path(__file__).parent.parent))

from models.rule import Rule, TransformationConfig
from utils.logger import get_logger


class RuleParser:
    """规则解析器"""

    def __init__(self, rules_path: str):
        self.rules_path = Path(rules_path)
        self.logger = get_logger()

        if not self.rules_path.exists():
            raise FileNotFoundError(f"规则文件不存在: {rules_path}")

    def parse(self) -> TransformationConfig:
        """
        解析规则文件

        Returns:
            TransformationConfig 对象
        """
        self.logger.step(1, "解析规则文件")
        self.logger.info(f"📄 读取规则文件: {self.rules_path}")

        # 打开 Excel 文件
        wb = openpyxl.load_workbook(self.rules_path, data_only=True)
        ws = wb.active

        rules = []
        field_order = []
        source_fields = set()

        # 数据从第7行开始（第6行是表头）
        self.logger.info("🔍 解析规则行...")
        row_count = 0

        for row in ws.iter_rows(min_row=7, max_row=ws.max_row, values_only=True):
            # 检查是否为空行
            if not row or len(row) < 2 or not row[1]:
                continue

            # 解析各列
            # 列索引: 0=CSDS_FLAG, 1=CS_COLUMN_NAME, 2=CARRIER_COLUMN_NAME,
            #         3=DEFAULT, 4=SPECIAL_RULES, 5=NOTE
            csds_flag = str(row[0]) if row[0] else ""
            cs_column = str(row[1]) if row[1] else ""
            carrier_column = str(row[2]) if row[2] else ""
            default = str(row[3]) if row[3] else ""
            special_rules = str(row[4]) if row[4] else ""
            note = str(row[5]) if len(row) > 5 and row[5] else ""

            # 跳过空字段名
            if not cs_column or cs_column in ["NULL", "nan"]:
                continue

            # 创建 Rule 对象
            rule = Rule(
                cs_column_name=cs_column,
                carrier_column_name=carrier_column
                if carrier_column not in ["NULL", "nan", ""]
                else "",
                default_value=default if default not in ["NULL", "nan", ""] else "",
                special_rules=special_rules
                if special_rules not in ["NULL", "nan", ""]
                else "",
                note=note,
            )

            rules.append(rule)
            field_order.append(cs_column)

            # 记录源字段
            if rule.carrier_column_name and rule.carrier_column_name not in [
                "NULL",
                "nan",
                "",
            ]:
                source_fields.add(rule.carrier_column_name)

            row_count += 1

        self.logger.success(
            f"解析完成: {row_count} 条规则, {len(source_fields)} 个源字段"
        )

        # 创建配置对象
        config = TransformationConfig(
            case_name=self.rules_path.stem,
            rules=rules,
            field_order=field_order,
            source_fields=sorted(list(source_fields)),
        )

        return config

    def parse_to_dict(self) -> Dict[str, Any]:
        """解析为字典格式"""
        config = self.parse()
        return config.to_dict()


if __name__ == "__main__":
    # 测试解析
    import sys

    sys.path.insert(0, str(Path(__file__).parent.parent.parent))

    # 测试 humanaS10
    parser = RuleParser("case/humanaS10/rules.xlsx")
    config = parser.parse()

    print(f"\n案例: {config.case_name}")
    print(f"规则数: {len(config.rules)}")
    print(f"字段数: {len(config.field_order)}")
    print(f"源字段: {config.source_fields[:10]}...")  # 前10个

    # 显示几个示例规则
    print("\n示例规则:")
    for rule in config.rules[:5]:
        print(f"  {rule}")
