"""
测试脚本：转换 highmark 案例数据

使用 utils 工具库转换 case/highmark/ 的源数据
"""

import sys
import importlib.util
from pathlib import Path
import openpyxl


# 加载工具模块
def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


base_path = Path(__file__).parent
shared_rules = load_module("shared_rules", base_path / "shared_rules.py")
file_readers = load_module("file_readers", base_path / "file_readers.py")
data_converters = load_module("data_converters", base_path / "data_converters.py")
formatters = load_module("formatters", base_path / "formatters.py")


def transform_highmark_data():
    """
    转换 highmark 案例的数据
    """
    print("=" * 80)
    print("测试：转换 highmark 案例数据")
    print("=" * 80)

    # 1. 读取源文件
    print("\n[1] 读取源文件...")
    source_df = file_readers.read_source_file("case/highmark/source.txt", delimiter="|")
    print(f"    ✅ 读取成功：{len(source_df)} 行, {len(source_df.columns)} 列")
    print(f"    列名: {list(source_df.columns[:8])}...")

    # 2. 读取 rules.xlsx 获取案例特定的固定值
    print("\n[2] 读取 rules.xlsx...")
    wb = openpyxl.load_workbook("case/highmark/rules.xlsx")
    ws = wb["Sheet1"]

    # 提取 CARRIER_FAMILY_ID
    carrier_family_id = None
    for row in ws.iter_rows(min_row=7, max_row=100):
        cs_col_name = row[1].value if len(row) > 1 else None
        source_col = row[3].value if len(row) > 3 else None

        if cs_col_name == "CARRIER_FAMILY_ID":
            carrier_family_id = data_converters.format_with_commas(source_col)
            break

        if not cs_col_name:
            break

    print(f"    ✅ CARRIER_FAMILY_ID: {carrier_family_id}")

    # 3. 提取字段映射（从 rules.xlsx）
    print("\n[3] 提取字段映射规则...")
    field_mappings = {}

    for row in ws.iter_rows(min_row=7, max_row=100):
        cs_col_name = row[1].value if len(row) > 1 else None
        carrier_col_name = (
            row[2].value if len(row) > 2 else None
        )  # 修正：读取 CARRIER_COLUMN_NAME (列C)

        if not cs_col_name:
            break

        # 跳过空字段和固定值字段
        if (
            carrier_col_name
            and carrier_col_name not in ["NULL", "n/a"]
            and not isinstance(carrier_col_name, (int, float))
            and carrier_col_name not in ["HOME", "1", "2"]
        ):
            field_mappings[cs_col_name] = str(carrier_col_name)

    print(f"    ✅ 找到 {len(field_mappings)} 个字段映射")
    print(f"    示例: {dict(list(field_mappings.items())[:3])}")

    # 4. 转换数据
    print("\n[4] 转换数据...")
    output_rows = []

    for idx, row in source_df.iterrows():
        output_row = {}

        # 应用共享固定值
        output_row.update(shared_rules.SHARED_FIXED_VALUES)

        # 应用案例特定固定值
        output_row["CARRIER_FAMILY_ID"] = carrier_family_id

        # 应用空字段
        for field in shared_rules.SHARED_EMPTY_FIELDS:
            output_row[field] = ""

        # 从源数据映射字段（基于 field_mappings）
        for target_field, source_field in field_mappings.items():
            value = data_converters.safe_get(row, source_field, "")

            # 特殊处理：日期字段
            if "DATE" in target_field and value:
                value = data_converters.convert_date_to_mmddyyyy(value)

            output_row[target_field] = value

        # LAST_TOUCHED_DATE 设为空
        output_row["LAST_TOUCHED_DATE"] = ""

        output_rows.append(output_row)

    print(f"    ✅ 转换完成：{len(output_rows)} 行")

    # 5. 生成输出
    print("\n[5] 生成输出文件...")
    metadata = formatters.create_metadata(
        action_id="highmark-cs-data-integration", service_map_id="10003358"
    )

    output_content = formatters.format_pipe_delimited_output(
        output_rows, shared_rules.EXPECTED_FIELD_ORDER, metadata
    )

    lines = output_content.split("\n")
    print(f"    ✅ 生成成功：共 {len(lines)} 行")

    # 显示输出预览
    print("\n    输出预览（前10行）:")
    for i, line in enumerate(lines[:10], 1):
        if len(line) > 120:
            print(f"      {i:2d}: {line[:120]}...")
        else:
            print(f"      {i:2d}: {line}")

    # 6. 保存文件
    output_file = "tmp/highmark_output.txt"
    Path("tmp").mkdir(exist_ok=True)
    with open(output_file, "w", encoding="utf-8") as f:
        f.write(output_content)
    print(f"\n    💾 输出已保存到: {output_file}")

    # 7. 对比验证
    print("\n[6] 对比验证...")

    # 读取期望输出
    with open("case/highmark/expected.txt", "r") as f:
        expected_lines = f.readlines()

    # 对比字段头
    output_header = lines[4].strip().split("|")
    expected_header = expected_lines[4].strip().split("|")

    print(f"    字段头对比:")
    print(f"      输出字段数: {len(output_header)}")
    print(f"      期望字段数: {len(expected_header)}")
    print(f"      字段头匹配: {'✅' if output_header == expected_header else '❌'}")

    # 对比 CARRIER_FAMILY_ID
    carrier_idx = output_header.index("CARRIER_FAMILY_ID")
    output_data = lines[5].strip().split("|")
    expected_data = expected_lines[5].strip().split("|")

    print(f"\n    CARRIER_FAMILY_ID 对比:")
    print(f"      输出: '{output_data[carrier_idx]}'")
    print(f"      期望: '{expected_data[carrier_idx]}'")
    print(
        f"      匹配: {'✅' if output_data[carrier_idx] == expected_data[carrier_idx] else '❌'}"
    )

    # 对比前3行数据的几个关键字段
    print(f"\n    前3行关键字段对比:")
    key_fields = [
        "CARRIER_FAMILY_ID",
        "IS_PAID",
        "BUSINESS_LINE",
        "FIRST_NAME",
        "LAST_NAME",
    ]

    for row_idx in range(min(3, len(output_rows))):
        output_data = lines[5 + row_idx].strip().split("|")
        expected_data = expected_lines[5 + row_idx].strip().split("|")

        print(f"\n      行{row_idx + 1}:")
        for field in key_fields:
            if field in output_header:
                field_idx = output_header.index(field)
                output_val = (
                    output_data[field_idx] if field_idx < len(output_data) else ""
                )
                expected_val = (
                    expected_data[field_idx] if field_idx < len(expected_data) else ""
                )
                match = "✅" if output_val == expected_val else "❌"
                print(
                    f"        {field:20s}: '{output_val}' vs '{expected_val}' {match}"
                )

    print("\n" + "=" * 80)
    print("✅ highmark 案例测试完成！")
    print("=" * 80)


if __name__ == "__main__":
    transform_highmark_data()
