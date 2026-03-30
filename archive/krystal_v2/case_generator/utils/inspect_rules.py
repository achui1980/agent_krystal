"""
Inspect rules.xlsx files from both humanaS10 and highmark cases
to understand the field mapping structure differences.
"""

import openpyxl
from pathlib import Path


def inspect_case_rules(case_name: str):
    """Inspect rules.xlsx for a specific case"""
    rules_path = Path(f"case/{case_name}/rules.xlsx")

    if not rules_path.exists():
        print(f"❌ Rules file not found: {rules_path}")
        return

    print(f"\n{'=' * 80}")
    print(f"📋 Inspecting: {case_name}/rules.xlsx")
    print(f"{'=' * 80}\n")

    wb = openpyxl.load_workbook(rules_path, data_only=True)
    ws = wb.active

    # Get header row (assuming row 6 based on previous knowledge)
    headers = []
    for col in range(1, ws.max_column + 1):
        header = ws.cell(row=6, column=col).value
        if header:
            headers.append(header)

    print(f"📊 Headers found: {headers}\n")

    # Find column indices
    col_indices = {}
    for idx, header in enumerate(headers, start=1):
        col_indices[header] = idx

    # Sample a few key fields to understand the pattern
    sample_fields = [
        "FIRST_NAME",
        "LAST_NAME",
        "BIRTH_DATE",
        "ADDRESS_LINE_1",
        "CARRIER_FAMILY_ID",
    ]

    print(f"🔍 Sampling key fields:\n")

    for row in ws.iter_rows(min_row=7, max_row=ws.max_row):
        cs_column = row[col_indices.get("CS_COLUMN_NAME", 2) - 1].value

        if cs_column in sample_fields:
            print(f"Field: {cs_column}")
            print(f"  Row Number: {row[0].row}")

            for header in headers:
                col_idx = col_indices[header] - 1
                value = row[col_idx].value
                print(f"  {header:20s}: {value}")
            print()

    # Summary statistics
    print(f"\n📈 Summary Statistics:")

    source_col_idx = col_indices.get("SOURCE_COLUMN", None)
    default_val_idx = col_indices.get("DEFAULT_VALUE", None)
    special_rules_idx = col_indices.get("SPECIAL_RULES", None)

    if source_col_idx:
        non_null_count = 0
        null_count = 0
        for row in ws.iter_rows(min_row=7, max_row=ws.max_row):
            val = row[source_col_idx - 1].value
            if val and str(val).strip().upper() != "NULL":
                non_null_count += 1
            else:
                null_count += 1

        print(f"  SOURCE_COLUMN: {non_null_count} non-NULL, {null_count} NULL")

    if default_val_idx:
        non_na_count = 0
        na_count = 0
        for row in ws.iter_rows(min_row=7, max_row=ws.max_row):
            val = row[default_val_idx - 1].value
            if val and str(val).strip().lower() != "n/a":
                non_na_count += 1
            else:
                na_count += 1

        print(f"  DEFAULT_VALUE: {non_na_count} non-n/a, {na_count} n/a")

    if special_rules_idx:
        non_empty_count = 0
        empty_count = 0
        for row in ws.iter_rows(min_row=7, max_row=ws.max_row):
            val = row[special_rules_idx - 1].value
            if val and str(val).strip():
                non_empty_count += 1
            else:
                empty_count += 1

        print(f"  SPECIAL_RULES: {non_empty_count} non-empty, {empty_count} empty")

    print(f"\n  Total rows (excluding header): {ws.max_row - 6}")

    wb.close()


def compare_specific_field(field_name: str, case1: str, case2: str):
    """Compare a specific field between two cases"""
    print(f"\n{'=' * 80}")
    print(f"🔬 Comparing field '{field_name}' between {case1} and {case2}")
    print(f"{'=' * 80}\n")

    for case in [case1, case2]:
        rules_path = Path(f"case/{case}/rules.xlsx")
        wb = openpyxl.load_workbook(rules_path, data_only=True)
        ws = wb.active

        # Get headers
        headers = []
        for col in range(1, ws.max_column + 1):
            header = ws.cell(row=6, column=col).value
            if header:
                headers.append((col, header))

        # Find the field
        for row in ws.iter_rows(min_row=7, max_row=ws.max_row):
            cs_column = row[1].value  # Assuming CS_COLUMN_NAME is column B (index 1)

            if cs_column == field_name:
                print(f"[{case}] {field_name}:")
                for col_idx, header in headers:
                    value = ws.cell(row=row[0].row, column=col_idx).value
                    print(f"  {header:20s}: {value}")
                print()
                break

        wb.close()


if __name__ == "__main__":
    # Inspect both cases
    inspect_case_rules("humanaS10")
    inspect_case_rules("highmark")

    # Compare specific fields
    compare_specific_field("FIRST_NAME", "humanaS10", "highmark")
    compare_specific_field("CARRIER_FAMILY_ID", "humanaS10", "highmark")
